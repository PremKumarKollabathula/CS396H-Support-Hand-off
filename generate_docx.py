from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START

OUTPUT_PATH = "CS396H-SOP-Word.docx"

DOC = {
    "title": "CS396H Support Hand-Off Documentation",
    "subtitle": "Vendor Export Auto Invoicing / EPIC EDD Override",
    "document_info": [
        ("Project Name", "CS396H - Vendor Export Auto Invoicing (EPIC EDD Override for 3PDS Vendors)"),
        ("Change Request", "Task 3071 - 3PDS Vendors / EDD Override for Epic-Enabled Warehouses"),
        ("Jira ID", "OMS-2246 (parent EDD override integration: OMS-2243)"),
        ("Revision Level Covered", "PK-L (single consolidated revision - approved via code review)"),
        ("Related Program", "CS396A (original/reference implementation)"),
        ("Developed by", "Prem Kumar K"),
        ("Date Created", "09/15/2026"),
        ("Support Hand-Off Date", "[Fill in]"),
        ("Document Version", "1.0"),
        ("Support Team", "Order Management - Primary Support"),
        ("Escalation Team", "IBM i Development Team / Nextuple EDD API Team"),
    ],
    "sections": [
        ("1. Project Scope", [
            ("1.1 Project Overview and Business Objectives", [
                "Program Name: CS396H - Build Vendor Export Auto Invoicing Records, enhanced to call the Nextuple EPIC EDD service before committing local ship/delivery date overrides for 3PDS orders.",
                "Business Purpose: CS396H processes queued vendor export transactions (VNDEXPQ, VNDEXPH, VNDEXPD) to build Auto-Invoicing records. For 3PDS orders shipping from an Epic-enabled warehouse in Live mode, the program calls the EPIC EDD service before finalizing local ship or delivery date overrides.",
                "Key Business Value:",
                "• Accurate dates: locally stored ship/delivery dates align with EPIC-calculated dates for Epic-enabled warehouses.",
                "• Per-item correction: date overrides use ExtOrd.RqsDat when it is later than the vendor/EPIC-derived date, applied consistently across all order line items.",
                "• Reliable requeue: genuine EPIC failures are retried instead of silently archiving or deleting the queue record as though the process succeeded.",
            ]),
            ("1.2 Key Enhancement - Revision PK-L", [
                "The approved current behavior under revision PK-L is:",
                "1. EDD_Override_SR (OMS-2243) checks isWarehouseEpicEnabled(Wk_House : IsWhseEnabled : EnablementMode). Only EnablementMode = 'L' (Live) is treated as Epic-enabled; blank or Compare mode is treated differently.",
                "2. Upd_Dates_Sr (OMS-2246) compares Wk_RqsDat (Packed 8:0) to ExtOrd.RqsDat against vedShpDt per line item. Every item on the order is checked independently.",
                "3. Requeue handling in Process_Rcds: when EPIC fails (IsWhseEnabled='Y' and EDDProcResponse.ProcessStatus <> 'PASS'), Continue_Flg is set to 'N', ensuring that the outer driver requeues the record instead of treating it as success.",
            ]),
        ]),
        ("2. Program Flow", [
            ("2.1 High-Level Flow", [
                "START -> Read VNDEXPQ where veqSts = Q",
                "      -> Chain VndExpQ / VndExpQA",
                "      -> ExSR Process_Rcds",
                "      -> ExSR Valid_SR (order validity + warehouse validation)",
                "      -> Continue_Flg = Y?",
                "          NO -> Hard_Err_Flg = 1; Continue_Flg = N",
                "          YES -> ExSR EDD_Override_SR (PK-L)",
                "                  -> isWarehouseEpicEnabled?",
                "                      NO -> ExSR Upd_Dates_Sr (non-EPIC local update)",
                "                      YES -> Loop through order items and call EPIC EDD service",
                "                             -> ProcessStatus = PASS?",
                "                                  YES -> ExSR Upd_Dates_Sr (apply EPIC-confirmed dates)",
                "                                  NO -> Continue_Flg = N",
                "      -> Outer driver checks Continue_Flg",
                "          YES -> Build transaction log, write archive queue, delete VNDEXPQR",
                "          NO -> Set veqSts = Q or E; update VNDEXPQR; write transaction log",
                "END",
            ]),
            ("2.2 Upd_Dates_Sr Detail", [
                "Loop CoDatNA1 by (copCusOrdN : vedItmNo)",
                "  Save WK_VedShpDt = vedShpDt",
                "  Chain(n) ExtOrd by copCusOrdN",
                "  If RqsDat > vedShpDt (as *ISO0), bump vedShpDt = RqsDat",
                "  If Hld_Order <> copCusOrdN (first item of order only): update CoMast (CO_RqDte / CO_MSDte)",
                "  Update CoDataCN / EXTORIT with corrected date",
                "This subroutine runs in both flows: directly for non-EPIC orders and after a successful EPIC response for Epic-enabled orders.",
            ]),
        ]),
        ("3. Support Responsibilities", [
            ("3.1 Primary Support Team: Order Management", [
                "Responsibilities:",
                "1. Monitor VNDEXPQ for records stuck at veqSts='Q' beyond expected retry windows.",
                "2. Monitor VNDEXPQ for records with veqSts='E' (hard error; will not auto-retry).",
                "3. Confirm whether a stuck order is due to EPIC EDD API failures or a hard data issue.",
                "4. Escalate genuine EPIC API failures to the Nextuple EDD API team.",
                "5. Escalate hard errors (Hard_Err_Flg='1'/'2'/'3') to the IBM i Development Team.",
                "What Support Does Not Do:",
                "• Modify CS396A/CS396H program logic.",
                "• Manually flip veqSts on VNDEXPQ without confirming the root cause.",
                "• Manually correct CoMast, CoDataN, or ExtOrIt dates without dev sign-off.",
            ]),
            ("3.2 Escalation Team: Nextuple EDD API Team", [
                "Role: Resolve EPIC EDD API errors or outages that prevent ProcessStatus from returning 'PASS' for Epic-enabled warehouse orders.",
                "Escalation required when:",
                "• Multiple VNDEXPQ records remain at veqSts='Q' for the same Epic-enabled warehouse across several job runs.",
                "• EDDProcResponse.ProcessStatus is consistently non-PASS for that warehouse when checked against the relevant health-monitor evidence.",
            ]),
        ]),
        ("4. VNDEXPQ Status Reference", [
            ("4.1 veqSts Values", [
                "Value / Meaning / Support Action:",
                "Q / Queued / requeued; will be picked up again on the next run / Normal after PK-L EPIC-fail requeue or a transient issue",
                "E / Hard error; will not auto-retry / Investigate before manual replay",
                "Row deleted / Success; archived to history / No action needed",
            ]),
            ("4.2 Distinguishing EPIC-Fail Requeue vs Hard Error", [
                "• EPIC-fail requeue: IsWhseEnabled='Y' and EDDProcResponse.ProcessStatus <> 'PASS'. Result: veqSts='Q', Continue_Flg='N', and a transaction log entry is written for the attempt.",
                "• Hard error: Hard_Err_Flg is set in Valid_SR when required records are missing or invalid. Result: veqSts='E'; no automatic requeue occurs.",
            ]),
        ]),
        ("5. Reprocessing / Retry Behavior", [
            ("5.1 On EPIC Failure", [
                "• IsWhseEnabled='Y' and ProcessStatus <> 'PASS' -> Continue_Flg='N' -> record is requeued (veqSts='Q') and a transaction log entry is written.",
                "• Non-EPIC orders (IsWhseEnabled='N') run Upd_Dates_Sr directly and are not affected by this logic.",
            ]),
            ("5.2 On Retry After a Prior EPIC Failure", [
                "• Each run of Process_Rcds re-evaluates the order independently.",
                "• If the retry succeeds (ProcessStatus='PASS'), Upd_Dates_Sr runs, Continue_Flg stays 'Y', and the outer driver logs, archives, and deletes the record as normal.",
                "• Support note: a transaction log is written for both failed and successful attempts. An order may therefore show multiple attempts before finally succeeding.",
            ]),
        ]),
        ("6. Support Procedures / SQL Queries", [
            ("Finding Stuck Requeued Records", [
                "SELECT * FROM VNDEXPQ WHERE VEQSTS = 'Q' ORDER BY <queue timestamp/key field>;",
            ]),
            ("Finding Hard-Error Records", [
                "SELECT * FROM VNDEXPQ WHERE VEQSTS = 'E' ORDER BY <queue timestamp/key field>;",
            ]),
            ("Cross-Reference with EPIC Health Monitor (CS448J)", [
                "Use the CS448J support hand-off and evidence queries against EDDREQHDR for the same warehouse and time window to confirm whether a stuck CS396H record is caused by a genuine EPIC API outage versus a local processing issue.",
            ]),
        ]),
        ("7. Program Dependencies", [
            ("Related / Reference Programs", [
                "• CS396A - original vendor export program with the reference EPIC EDD override implementation.",
                "• CS448H - EpicEnabledWarehouseList / isWarehouseEpicEnabled check.",
                "• CS448J - EDD Service Health Monitor; used to confirm genuine EPIC API outages.",
            ]),
            ("Key Tables", [
                "• VNDEXPQ, VNDEXPH, VNDEXPD - vendor export queue, header, and detail.",
                "• COMAST, CODATAN, EXTORIT, EXTORD - order master and item-level date fields.",
            ]),
        ]),
        ("8. Contact Information", [
            ("Support contacts", [
                "L1 / Order Management Support / [Support email/queue] / 30 min (P2)",
                "L2 / Nextuple EDD API Team / [Nextuple support contact] / 1 hour (P2)",
                "L3 / IBM i Development Team / [Dev team contact] / 2 hours (P2)",
                "Original / Enhancement Developer: Prem Kumar K",
            ]),
        ]),
        ("Document Approval", [
            "• Document Prepared By: Development Team",
            "• Reviewed By: [Name]",
            "• Approved By: [Name]",
            "• Date: [Fill in]",
        ]),
    ],
}



def add_title_page(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(DOC["title"])
    run.bold = True
    run.font.size = Pt(26)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = subtitle.add_run(DOC["subtitle"])
    run2.italic = True
    run2.font.size = Pt(16)

    doc.add_paragraph()
    doc.add_paragraph("Prepared for support handoff and operational continuity")
    doc.add_paragraph()
    doc.add_paragraph("Version 1.0")
    doc.add_page_break()


def add_document_info_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    for label, value in DOC["document_info"]:
        row_cells = table.add_row().cells
        row_cells[0].text = label
        row_cells[1].text = value

    doc.add_paragraph()


def add_section(doc, heading, content_items):
    h = doc.add_heading(heading, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for item in content_items:
        if isinstance(item, tuple):
            subheading, subitems = item
            s = doc.add_heading(subheading, level=2)
            s.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for text in subitems:
                if isinstance(text, list):
                    continue
                p = doc.add_paragraph(text)
                if text.startswith("•"):
                    p.style = "List Bullet"
                elif text.startswith("1.") or text.startswith("2.") or text.startswith("3."):
                    p.style = "List Number"
        else:
            p = doc.add_paragraph(item)
            if item.startswith("•"):
                p.style = "List Bullet"
            elif item.startswith("1.") or item.startswith("2.") or item.startswith("3."):
                p.style = "List Number"


def build_document():
    document = Document()
    add_title_page(document)
    add_document_info_table(document)

    for heading, sections in DOC["sections"]:
        document.add_heading(heading, level=1)
        for subheading, paragraphs in sections:
            document.add_heading(subheading, level=2)
            for para in paragraphs:
                p = document.add_paragraph(para)
                if para.startswith("•"):
                    p.style = "List Bullet"
                elif para.startswith("1.") or para.startswith("2.") or para.startswith("3."):
                    p.style = "List Number"

    doc = document
    doc.add_paragraph()
    doc.add_paragraph("END OF SUPPORT HAND-OFF DOCUMENTATION")
    doc.save(OUTPUT_PATH)
    print(f"Created {OUTPUT_PATH}")


if __name__ == "__main__":
    build_document()
